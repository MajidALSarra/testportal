#!/usr/bin/env python3
"""Build the reviewer demo PC's fictional documents into one zip.

Everything here is invented: the company "Blue Heron Outfitters", its people
and its numbers. No real person, customer or file is used. Run:
    python3 make_seed_files.py <out.zip>
The zip mirrors the target profile: Desktop/, Documents/, Documents/Reports/.
Photos are NOT in here; the VM downloads public-domain/CC0 photos itself
(see ../vm/seed-photos.ps1) because this container has no route to them.
"""
import io, random, sys, zipfile, datetime as dt
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table

CO = "Blue Heron Outfitters"
rnd = random.Random(2026)  # fixed seed: same numbers every build
MONTHS = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
PRODUCTS = ["Trail Pack 30L","Rain Shell","Camp Stove","Hiking Poles","Sleeping Bag","Headlamp"]
REGIONS = ["North","South","East","West"]

def pdf(title, paras, table=None):
    b = io.BytesIO(); s = getSampleStyleSheet()
    flow = [Paragraph(title, s["Title"]), Paragraph(CO, s["Heading3"]), Spacer(1, 12)]
    for p in paras:
        flow += [Paragraph(p, s["BodyText"]), Spacer(1, 8)]
    if table: flow.append(Table(table))
    SimpleDocTemplate(b, pagesize=A4, title=title, author=CO).build(flow)
    return b.getvalue()

def docx(title, paras):
    d = Document(); d.core_properties.author = CO; d.core_properties.title = title
    d.add_heading(title, 0)
    for p in paras: d.add_paragraph(p)
    b = io.BytesIO(); d.save(b); return b.getvalue()

def xlsx(sheets):
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows: ws.append(r)
        for c in ws[1]: c.font = Font(bold=True)
    wb.properties.creator = CO
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def pptx(title, slides):
    p = Presentation(); p.core_properties.author = CO
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = title; s.placeholders[1].text = CO
    for head, bullets in slides:
        s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = head
        tf = s.placeholders[1].text_frame; tf.text = bullets[0]
        for x in bullets[1:]: tf.add_paragraph().text = x
    b = io.BytesIO(); p.save(b); return b.getvalue()

def sales_rows():
    rows = [["Month","Region","Product","Units","Revenue (USD)"]]
    for m in ["Jul","Aug","Sep"]:
        for r in REGIONS:
            for pr in PRODUCTS:
                u = rnd.randint(20, 400); rows.append([m, r, pr, u, round(u * rnd.uniform(18, 140), 2)])
    return rows

def build():
    f = {}
    f["Documents/Annual-Report-2026.pdf"] = pdf("Annual Report 2026", [
        "This annual report summarises the year for Blue Heron Outfitters, a fictional outdoor-gear retailer used for product demonstrations.",
        "Revenue grew 12% year over year, led by the Trail Pack and Rain Shell lines. Online orders were 41% of sales.",
        "Next year we plan two new stores, a repair service and a refreshed loyalty programme."],
        [["Quarter","Revenue (USD k)","Growth"],["Q1","1,240","8%"],["Q2","1,410","11%"],["Q3","1,655","14%"],["Q4","1,820","13%"]])
    f["Documents/Reports/Monthly-Report-September.pdf"] = pdf("Monthly Report - September", [
        "September closed above plan. The West region led unit sales; returns stayed under 2%.",
        "Inventory of sleeping bags is low ahead of winter; a reorder is scheduled for October."])
    f["Documents/Employee-Handbook.pdf"] = pdf("Employee Handbook", [
        "Welcome to Blue Heron Outfitters. This handbook covers working hours, leave, safety on the shop floor and how to request equipment.",
        "Questions go to the (fictional) people team at people@blueheron.example."])
    f["Documents/Project-Report-Draft.docx"] = docx("Project Report (Draft)", [
        "Project: new online store checkout.", "Status: on track. Two of three milestones are done.",
        "Risks: payment provider testing may take an extra week.", "Next step: user testing with ten customers in November."])
    f["Documents/Meeting-Notes-2026-10.docx"] = docx("Meeting Notes - October 2026", [
        "Attendees: Sam Rivera, Priya Nair, Tom Okafor (all fictional).",
        "1. Winter campaign starts 1 November.", "2. Store opening hours unchanged.", "3. Next meeting: 3 November."])
    f["Documents/Q3-Sales-Report.xlsx"] = xlsx([("Q3 Sales", sales_rows())])
    budget = [["Category"] + MONTHS] + [[c] + [rnd.randint(5, 60) * 1000 for _ in MONTHS]
              for c in ["Rent","Salaries","Marketing","Stock","Shipping","Software"]]
    f["Documents/Budget-2027.xlsx"] = xlsx([("Budget 2027", budget)])
    f["Documents/Report-Template.pptx"] = pptx("Report Template", [
        ("Summary", ["Key result 1", "Key result 2", "Key result 3"]),
        ("Numbers", ["Revenue", "Customers", "Returns"]), ("Next steps", ["Owner and date for each action"])])
    f["Documents/Product-Roadmap.pptx"] = pptx("Product Roadmap 2027", [
        ("Spring", ["Ultralight tent", "Kids' rain shell"]), ("Summer", ["Solar lantern", "Trail Pack 45L"]),
        ("Autumn", ["Repair service in all stores"])])
    f["Desktop/notes.txt"] = "Demo notes\n- Order more headlamps\n- Call the printer about flyers\n- Team lunch on Friday\n"
    f["Documents/notes.txt"] = "Shopping list for the store opening: banners, ribbon, coffee, name badges.\n"
    f["Desktop/Welcome.txt"] = ("This is a demo computer for testing Folderbird.\n"
        "All files here are fictional sample files. Photos are public-domain or CC0 images.\n")
    return f

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "seed-docs.zip"
    stamp = (2026, 10, 1, 9, 0, 0)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in sorted(build().items()):
            zi = zipfile.ZipInfo(name, stamp); zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data if isinstance(data, bytes) else data.encode())
    print(out)
